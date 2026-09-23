"""关键节点（上游近似最小割点）算法与接口测试。

运行：python -m pytest tests/test_keynodes.py -v
"""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.graph import PATH_COUNT_CAP, WorkbookGraph
from app.main import app

client = TestClient(app)


# ---------------------------------------------------------------- 构图工具
def build_graph(wb: Workbook) -> WorkbookGraph:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wbf = load_workbook(buf, data_only=False)
    buf.seek(0)
    wbv = load_workbook(buf, data_only=True)
    return WorkbookGraph.from_workbooks(wbf, wbv)


def diamond_wb() -> Workbook:
    """两个独立输入 -> 汇聚点 C1 -> 两条分支 C2/C3 -> 目标 D1。"""
    wb = Workbook()
    p = wb.active
    p.title = "参数"
    p["A1"] = 1
    p["B1"] = 2
    c = wb.create_sheet("计算")
    c["C1"] = "=参数!A1+参数!B1"
    c["C2"] = "=C1*2"
    c["C3"] = "=C1+1"
    s = wb.create_sheet("汇总")
    s["D1"] = "=计算!C2+计算!C3"
    return wb


def by_cell(res: dict) -> dict:
    return {r["cell"]: r for r in res["key_nodes"]}


# ---------------------------------------------------------------- 算法
def test_coverage_counts_diamond():
    g = build_graph(diamond_wb())
    res = g.key_nodes("汇总!D1")
    rows = by_cell(res)
    # 汇聚点 C1：入路径 2（A1/B1）× 出路径 2（经 C2/C3）= 4，覆盖最多
    assert res["key_nodes"][0]["cell"] == "计算!C1"
    assert rows["计算!C1"]["coverage"] == 4
    assert rows["计算!C1"]["in_paths"] == 2
    assert rows["计算!C1"]["out_paths"] == 2
    # 输入叶子只有一条路径经过；C2/C3 各自只走一条出路径
    assert rows["参数!A1"]["coverage"] == 1 * 2
    assert rows["计算!C2"]["coverage"] == 2 * 1
    # 目标自身不计入候选
    assert "汇总!D1" not in rows
    assert res["source_total"] == 2
    assert res["upstream_total"] == 5
    assert res["collapsed_scc_count"] == 0


def test_rank_order_and_fields():
    g = build_graph(diamond_wb())
    res = g.key_nodes("汇总!D1")
    covs = [r["coverage"] for r in res["key_nodes"]]
    assert covs == sorted(covs, reverse=True)  # 默认按覆盖数降序
    assert [r["rank"] for r in res["key_nodes"]] == list(range(1, len(covs) + 1))
    top = res["key_nodes"][0]
    assert {"cell", "sheet", "display", "coverage", "in_paths", "out_paths",
            "affected_downstream", "affected_downstream_count",
            "evidence_path", "evidence_path_raw", "component_members"} <= set(top)


def test_affected_downstream_and_evidence():
    g = build_graph(diamond_wb())
    res = g.key_nodes("汇总!D1")
    rows = by_cell(res)
    # 移除 C1 会切断 C2、C3 到目标的路径（目标本身不计入）
    aff = {a["cell"] for a in rows["计算!C1"]["affected_downstream"]}
    assert aff == {"计算!C2", "计算!C3"}
    # C2 直接流向目标，无额外受影响下游
    assert rows["计算!C2"]["affected_downstream_count"] == 0
    # 证据路径：源 -> 该节点 -> 目标，且节点自身必在路径中
    ev = rows["参数!A1"]["evidence_path"]
    assert ev[0] == "参数!A1" and ev[-1] == "汇总!D1"
    assert "计算!C1" in ev
    raw = rows["参数!A1"]["evidence_path_raw"]
    assert raw[0] == "参数!A1" and raw[-1] == "汇总!D1"


def test_evidence_paths_are_real_graph_walks():
    g = build_graph(diamond_wb())
    res = g.key_nodes("汇总!D1")
    # 每条证据路径的相邻单元格在依赖图中必须有真实边（dependent 方向）
    for r in res["key_nodes"]:
        path = r["evidence_path_raw"]
        for a, b in zip(path, path[1:]):
            assert b in g.dependents.get(a, set()), (a, b, path)


def test_no_upstream_is_empty():
    g = build_graph(diamond_wb())
    res = g.key_nodes("参数!A1")
    assert res["key_nodes"] == []
    assert res["total"] == 0 and res["upstream_total"] == 0


def test_cycle_scc_is_collapsed():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["C1"] = 5
    ws["A1"] = "=B1+1"
    ws["B1"] = "=A1+C1"   # A1 <-> B1 构成 SCC，C1 为外部输入
    ws["D1"] = "=A1*2"    # 目标
    g = build_graph(wb)

    res = g.key_nodes("S!D1")
    assert res["collapsed_scc_count"] == 1
    # SCC 压缩为单个候选节点，代表为排序首个成员 A1
    scc_rows = [r for r in res["key_nodes"] if r["is_component"]]
    assert len(scc_rows) == 1
    scc = scc_rows[0]
    assert scc["cell"] == "S!A1"
    assert {m["cell"] for m in scc["component_members"]} == {"S!A1", "S!B1"}
    # SCC 覆盖唯一一条 源(C1)->目标(D1) 路径
    assert scc["coverage"] == 1
    # 证据路径穿过 SCC：C1 -> (B1/A1) -> D1
    raw = scc["evidence_path_raw"]
    assert raw[0] == "S!C1" and raw[-1] == "S!D1"
    assert {"S!A1", "S!B1"} <= set(raw)
    # 普通节点不标记为压缩分量
    leaf = by_cell(res)["S!C1"]
    assert leaf["is_component"] is False
    assert len(leaf["component_members"]) == 1


def test_self_loop_scc():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["B1"] = "=B1+A1"   # 自环 + 依赖 A1
    ws["C1"] = "=B1"      # 目标
    g = build_graph(wb)
    res = g.key_nodes("S!C1")
    rows = by_cell(res)
    # 自环 SCC 只有一个成员，压缩后仍是单节点；以 in_cycle 标记循环身份
    assert "S!B1" in rows
    assert rows["S!B1"]["in_cycle"] is True
    assert rows["S!B1"]["is_component"] is False


def test_parallel_paths_all_counted():
    """3 条完全独立的输入直达目标：每个输入覆盖 1 条路径，无瓶颈。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = 3
    ws["B1"] = "=A1+A2+A3"  # 目标
    g = build_graph(wb)
    res = g.key_nodes("S!B1")
    rows = by_cell(res)
    assert all(rows[f"S!A{i}"]["coverage"] == 1 for i in (1, 2, 3))
    assert res["source_total"] == 3


def test_path_count_saturation(monkeypatch):
    """多层完全二分菱形使路径数指数增长；调小上限后路径数截断并标记 saturated。"""
    from openpyxl.utils import get_column_letter
    import app.graph as graph_mod

    monkeypatch.setattr(graph_mod, "PATH_COUNT_CAP", 1000)

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    # 每层 2 个节点，均同时引用上一层全部 2 个节点：路径数每层翻倍
    prev = ["A1"]
    for layer in range(1, 13):  # 12 层后路径数 2048 > 1000
        cur = [f"{get_column_letter(2 * layer)}1",
               f"{get_column_letter(2 * layer + 1)}1"]
        for tgt in cur:
            ws[tgt] = "=" + "+".join(prev)
        prev = cur
    target = "A2"  # 第 2 行，避免与分层节点（都在第 1 行）坐标冲突
    ws[target] = "=" + "+".join(prev)
    g = build_graph(wb)
    res = g.key_nodes(f"S!{target}")
    assert res["path_count_cap"] == 1000
    assert any(r["saturated"] for r in res["key_nodes"])
    assert all(r["coverage"] <= 1000 for r in res["key_nodes"])


# ---------------------------------------------------------------- API
def upload(wb: Workbook) -> str:
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(
        "/api/upload",
        files={"file": ("k.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    return r.json()["workbook_id"]


def test_key_nodes_endpoint():
    wid = upload(diamond_wb())
    r = client.get(f"/api/workbooks/{wid}/key-nodes",
                   params={"cell": "汇总!D1"})
    assert r.status_code == 200
    d = r.json()
    assert d["root"] == "汇总!D1"
    assert d["key_nodes"][0]["cell"] == "计算!C1"
    assert d["key_nodes"][0]["coverage"] == 4
    assert d["total"] == 5 and len(d["key_nodes"]) == 5
    assert d["filters"] == {"sheet": None, "sort": "coverage", "min_coverage": 1}


def test_key_nodes_quoted_display_and_sheet_filter():
    wid = upload(diamond_wb())
    # 展示形式（带引号）也能解析
    r = client.get(f"/api/workbooks/{wid}/key-nodes",
                   params={"cell": "'汇总'!D1", "sheet": "参数"})
    assert r.status_code == 200
    d = r.json()
    assert {k["cell"] for k in d["key_nodes"]} == {"参数!A1", "参数!B1"}
    assert all(k["sheet"] == "参数" for k in d["key_nodes"])


def test_key_nodes_sorting_and_paging():
    wid = upload(diamond_wb())
    r = client.get(f"/api/workbooks/{wid}/key-nodes",
                   params={"cell": "汇总!D1", "sort": "cell", "limit": 2})
    d = r.json()
    assert d["total"] == 5 and len(d["key_nodes"]) == 2
    assert [k["cell"] for k in d["key_nodes"]] == ["参数!A1", "参数!B1"]
    r = client.get(f"/api/workbooks/{wid}/key-nodes",
                   params={"cell": "汇总!D1", "sort": "affected", "limit": 5})
    affs = [k["affected_downstream_count"] for k in r.json()["key_nodes"]]
    assert affs == sorted(affs, reverse=True)
    # min_coverage 过滤
    r = client.get(f"/api/workbooks/{wid}/key-nodes",
                   params={"cell": "汇总!D1", "min_coverage": 4})
    assert [k["cell"] for k in r.json()["key_nodes"]] == ["计算!C1"]


def test_key_nodes_export_is_untruncated():
    wid = upload(diamond_wb())
    r = client.get(f"/api/workbooks/{wid}/key-nodes/export",
                   params={"cell": "汇总!D1"})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert "key_nodes_" in r.headers["content-disposition"]
    d = r.json()
    assert d["root"] == "汇总!D1"
    assert d["total"] == 5
    # 导出中受影响下游完整（该例中仅 2 个，未达导出上限）
    top = next(k for k in d["key_nodes"] if k["cell"] == "计算!C1")
    assert len(top["affected_downstream"]) == top["affected_downstream_count"]
    assert "evidence_path_raw" in top
    assert "affected_truncated" not in top


def test_key_nodes_affected_truncation():
    """每个节点的受影响下游在接口上限 100、导出上限 1000 处截断并带标记。"""
    from app.graph import MAX_KEYNODE_AFFECTED, MAX_KEYNODE_AFFECTED_EXPORT

    m = MAX_KEYNODE_AFFECTED + 20  # 120 个 H1 的下游中转
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["H1"] = 1
    for i in range(m):
        ws.cell(row=i + 2, column=3, value=f"=H1+{i}")   # C2.. 全部依赖 H1
    ws["T1"] = "=" + "+".join(f"C{i+2}" for i in range(m))  # 目标汇聚全部中转
    wid = upload(wb)

    r = client.get(f"/api/workbooks/{wid}/key-nodes", params={"cell": "S!T1"})
    h1 = next(k for k in r.json()["key_nodes"] if k["cell"] == "S!H1")
    assert h1["affected_downstream_count"] == m
    assert len(h1["affected_downstream"]) == MAX_KEYNODE_AFFECTED
    assert h1["affected_truncated"] is True

    # 导出上限更高（1000），120 个下游可完整返回
    r = client.get(f"/api/workbooks/{wid}/key-nodes/export",
                   params={"cell": "S!T1"})
    h1 = next(k for k in r.json()["key_nodes"] if k["cell"] == "S!H1")
    assert len(h1["affected_downstream"]) == m
    assert "affected_truncated" not in h1

    # 超过导出上限时同样截断
    big = Workbook()
    bws = big.active
    bws.title = "S"
    bws["H1"] = 1
    mm = MAX_KEYNODE_AFFECTED_EXPORT + 10
    for i in range(mm):
        bws.cell(row=i + 2, column=3, value=f"=H1+{i}")
    bws["T1"] = "=" + "+".join(f"C{i+2}" for i in range(mm))
    wid2 = upload(big)
    r = client.get(f"/api/workbooks/{wid2}/key-nodes/export",
                   params={"cell": "S!T1"})
    h1 = next(k for k in r.json()["key_nodes"] if k["cell"] == "S!H1")
    assert h1["affected_downstream_count"] == mm
    assert len(h1["affected_downstream"]) == MAX_KEYNODE_AFFECTED_EXPORT
    assert h1["affected_truncated"] is True


def test_key_nodes_bad_params():
    wid = upload(diamond_wb())
    assert client.get(f"/api/workbooks/{wid}/key-nodes",
                      params={"cell": "D1"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/key-nodes",
                      params={"cell": "幽灵!A1"}).status_code == 404
    assert client.get(f"/api/workbooks/{wid}/key-nodes",
                      params={"cell": "汇总!D1", "sheet": "不存在"}).status_code == 404
    assert client.get(f"/api/workbooks/{wid}/key-nodes",
                      params={"cell": "汇总!D1", "sort": "bogus"}).status_code == 422
    assert client.get(f"/api/workbooks/{wid}/key-nodes",
                      params={"cell": "汇总!D1", "min_coverage": 0}).status_code in (400, 422)
    assert client.get("/api/workbooks/deadbeefab/key-nodes",
                      params={"cell": "汇总!D1"}).status_code == 404
    assert client.get("/api/workbooks/deadbeefab/key-nodes/export",
                      params={"cell": "汇总!D1"}).status_code == 404
