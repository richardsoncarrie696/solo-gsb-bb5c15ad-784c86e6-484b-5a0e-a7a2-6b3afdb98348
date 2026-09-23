"""关键节点（上游近似最小割点）分析测试。运行：python -m pytest tests/ -v"""
import io

from openpyxl import Workbook, load_workbook

from app.graph import WorkbookGraph


def _graph(wb: Workbook) -> WorkbookGraph:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wbf = load_workbook(buf, data_only=False)
    buf.seek(0)
    wbv = load_workbook(buf, data_only=True)
    return WorkbookGraph.from_workbooks(wbf, wbv)


def diamond_graph() -> WorkbookGraph:
    """
    A1 ─┐
        B1 ─┐
    A2 ─┘   │
        B2 ─┴─ C1 ─ D1
    A3 ─┘
    源→目标(D1)路径共 4 条：A1→B1→C1→D1、A2→B1→C1→D1、
    A2→B2→C1→D1、A3→B2→C1→D1；C1 是唯一必经单点。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
    ws["B1"] = "=A1+A2"
    ws["B2"] = "=A2+A3"
    ws["C1"] = "=B1+B2"
    ws["D1"] = "=C1*100"
    return _graph(wb)


def test_path_counts_and_cut_node():
    g = diamond_graph()
    r = g.key_nodes("S1!D1")
    assert r["total_paths"] == 4
    by = {k["cell"]: k for k in r["key_nodes"]}
    # C1 覆盖全部 4 条路径：必经节点
    assert by["S1!C1"]["covered_paths"] == 4
    assert by["S1!C1"]["cuts_all_paths"] is True
    assert by["S1!C1"]["coverage_ratio"] == 1.0
    # A2 是 B1、B2 的公共上游，覆盖 2 条；A1/A3 各覆盖 1 条
    assert by["S1!A2"]["covered_paths"] == 2
    assert by["S1!A1"]["covered_paths"] == 1
    assert by["S1!A3"]["covered_paths"] == 1
    assert by["S1!A1"]["cuts_all_paths"] is False
    # 默认按覆盖数降序
    covs = [k["covered_paths"] for k in r["key_nodes"]]
    assert covs == sorted(covs, reverse=True)
    assert r["key_nodes"][0]["cell"] == "S1!C1"


def test_affected_downstream_includes_target():
    g = diamond_graph()
    r = g.key_nodes("S1!D1")
    by = {k["cell"]: k for k in r["key_nodes"]}
    # 移除 C1 后目标 D1 无法计算
    assert by["S1!C1"]["affected_downstream"] == ["S1!D1"]
    assert by["S1!C1"]["affected_downstream_count"] == 1
    # B1 的下游：C1、D1
    assert set(by["S1!B1"]["affected_downstream"]) == {"S1!C1", "S1!D1"}
    # A1 的下游：B1、C1、D1
    assert by["S1!A1"]["affected_downstream_count"] == 3


def test_evidence_path_passes_through_node():
    g = diamond_graph()
    r = g.key_nodes("S1!D1")
    for k in r["key_nodes"]:
        assert k["evidence_path_raw"][0].startswith("S1!A")  # 起于输入源
        assert k["evidence_path_raw"][-1] == "S1!D1"         # 终于目标
        assert k["cell"] in k["evidence_path_raw"]           # 经过节点自身
        # 证据路径相邻分量在压缩 DAG 上确实有边
        for a, b in zip(k["evidence_path_raw"], k["evidence_path_raw"][1:]):
            assert b in g.dependents[a]


def test_sheet_filter_and_sorts():
    g = diamond_graph()
    r = g.key_nodes("S1!D1", sheet="S1")
    # A1/A2/A3/B1/B2/C1 六个候选（目标 D1 自身排除）
    assert len(r["key_nodes"]) == 6
    assert all(k["sheet"] == "S1" for k in r["key_nodes"])

    by_depth = g.key_nodes("S1!D1", sort="depth")["key_nodes"]
    cells = [k["cell"] for k in by_depth]
    # A1/A2/A3 同属第 3 层（按单元格名排序），B1/B2 第 2 层，C1 第 1 层
    assert set(cells[:3]) == {"S1!A1", "S1!A2", "S1!A3"}
    assert cells[3:5] == ["S1!B1", "S1!B2"]
    assert cells[5] == "S1!C1"
    depths = [k["depth"] for k in by_depth]
    assert depths == sorted(depths, reverse=True)
    assert by_depth[0]["depth"] == 3 and by_depth[-1]["depth"] == 1

    by_aff = g.key_nodes("S1!D1", sort="affected")["key_nodes"]
    # A2 同时通向 B1、B2，下游（B1/B2/C1/D1）最多
    assert by_aff[0]["cell"] == "S1!A2"
    assert by_aff[0]["affected_downstream_count"] == 4

    r = g.key_nodes("S1!D1", min_coverage=2)
    # C1 覆盖 4 条；A2/B1/B2 各覆盖 2 条；A1/A3 仅 1 条
    assert {k["cell"] for k in r["key_nodes"]} == {
        "S1!A2", "S1!B1", "S1!B2", "S1!C1"}
    r = g.key_nodes("S1!D1", min_coverage=4)
    assert {k["cell"] for k in r["key_nodes"]} == {"S1!C1"}


def test_leaf_target_has_no_key_nodes():
    g = diamond_graph()
    r = g.key_nodes("S1!A1")
    assert r["total_paths"] == 0
    assert r["key_nodes"] == []
    assert r["source_count"] == 0


def cycle_graph() -> WorkbookGraph:
    """A1↔B1 构成 SCC；C1 自环；D1 依赖 A1（环外下游）；E1 为目标。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "环"
    ws["A1"] = 1
    ws["B1"] = "=A1*2"
    ws["A1"] = "=B1+1"   # 覆盖原值，A1↔B1 互引
    ws["C1"] = "=C1*2"   # 自环 SCC
    ws["D1"] = "=A1+5"
    ws["E1"] = "=D1+C1"
    return _graph(wb)


def test_scc_collapsed_to_single_node():
    g = cycle_graph()
    r = g.key_nodes("环!E1")
    scc_rows = [k for k in r["key_nodes"] if k["type"] == "scc"]
    # {A1,B1} 与 {C1} 两个循环分量各压缩为一个候选节点
    member_sets = {tuple(sorted(k["members"])) for k in scc_rows}
    assert ("环!A1", "环!B1") in member_sets
    assert ("环!C1",) in member_sets
    for k in scc_rows:
        assert k["in_cycle"] is True
        # E1 有两条独立上游支链（经 D1 / 经 C1），单分量覆盖 1/2，非全局割点
        assert k["covered_paths"] == 1
        assert k["cuts_all_paths"] is False
    # 证据路径把 SCC 作为一个整体步骤展示
    ab = next(k for k in scc_rows if len(k["members"]) == 2)
    comp_types = [c["type"] for c in ab["evidence_components"]]
    assert "scc" in comp_types
    # SCC 行没有单一公式 / 值
    assert ab["formula"] is None and ab["cell"] is None


def test_scc_on_unique_chain_is_cut_node():
    """循环位于唯一的上游链路时，压缩后的 SCC 即必经节点。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "环2"
    ws["X1"] = 5
    ws["A1"] = "=B1+X1"
    ws["B1"] = "=A1*2"   # A1↔B1 SCC，是目标 D2 唯一的上游链
    ws["D2"] = "=A1"
    g = _graph(wb)
    r = g.key_nodes("环2!D2")
    assert r["total_paths"] == 1
    scc = next(k for k in r["key_nodes"] if k["type"] == "scc")
    assert tuple(sorted(scc["members"])) == ("环2!A1", "环2!B1")
    assert scc["covered_paths"] == 1 and scc["cuts_all_paths"] is True
    # 移除该 SCC 后受影响的下游含目标 D2
    assert "环2!D2" in scc["affected_downstream"]


def test_self_loop_scc_is_cyclic():
    g = cycle_graph()
    r = g.key_nodes("环!E1")
    c1 = next(k for k in r["key_nodes"] if k["members"] == ["环!C1"])
    assert c1["type"] == "scc" and c1["in_cycle"] is True


def parallel_chain_graph() -> WorkbookGraph:
    """两条完全独立的支链汇入目标，验证无单点割点时 cuts_all 均为 False。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "P"
    ws["A1"], ws["C1"] = 1, 2
    ws["B1"] = "=A1"
    ws["D1"] = "=C1"
    ws["E1"] = "=B1+D1"
    return _graph(wb)


def test_parallel_chains_have_no_single_cut():
    g = parallel_chain_graph()
    r = g.key_nodes("P!E1")
    assert r["total_paths"] == 2
    assert all(not k["cuts_all_paths"] for k in r["key_nodes"])
    assert all(k["covered_paths"] == 1 for k in r["key_nodes"])


def test_export_rows_carry_all_fields():
    g = diamond_graph()
    r = g.key_nodes("S1!D1")
    required = {
        "cell", "type", "display", "members", "in_cycle", "depth",
        "covered_paths", "covered_paths_text", "coverage_ratio",
        "cuts_all_paths", "affected_downstream_count", "affected_downstream",
        "affected_downstream_display", "evidence_path", "evidence_path_raw",
        "evidence_components",
    }
    for k in r["key_nodes"]:
        assert required <= set(k)
    assert isinstance(r["total_paths"], int)
    assert r["sampled_paths"] == r["total_paths"]
    assert r["path_truncated"] is False


def test_export_unlimited_returns_full_downstream():
    """limit=None 时受影响下游样本不截断（导出场景）。"""
    g = parallel_chain_graph()
    r = g.key_nodes("P!E1", limit=None)
    assert r["returned"] == r["total"]
    for k in r["key_nodes"]:
        assert len(k["affected_downstream"]) == k["affected_downstream_count"]
